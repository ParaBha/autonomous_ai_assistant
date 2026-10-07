import fitz  # PyMuPDF
import docx
from pptx import Presentation
import pandas as pd
import os
import io
from app.core.config import settings

class IngestionService:
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        text = ""
        filename = os.path.basename(file_path)
        try:
            with fitz.open(file_path) as doc:
                if doc.is_encrypted:
                    try:
                        doc.authenticate("")
                    except Exception:
                        return f"[PDF Document Asset: {filename} (Password Protected or Encrypted PDF)]"
                for page in doc:
                    page_text = page.get_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            return f"[PDF Document Asset: {filename} (Extraction note: {str(e)})]"

        # If text is empty (scanned image-only PDF), attempt OCR page rendering
        if not text.strip():
            try:
                with fitz.open(file_path) as doc:
                    ocr_pages = []
                    for page_num, page in enumerate(doc, start=1):
                        pix = page.get_pixmap(dpi=150)
                        from PIL import Image
                        import pytesseract
                        img = Image.open(io.BytesIO(pix.tobytes("png")))
                        p_ocr = pytesseract.image_to_string(img).strip()
                        if p_ocr:
                            ocr_pages.append(f"--- PAGE {page_num} (OCR) ---\n{p_ocr}")
                    if ocr_pages:
                        return "\n\n".join(ocr_pages)
            except Exception as ocr_err:
                print(f"PDF OCR fallback notice for {filename}: {ocr_err}")
            return f"[PDF Document Asset: {filename} (Scanned Image-Based PDF)]"

        return text

    @staticmethod
    def extract_text_with_pages(file_path: str) -> str:
        ext = file_path.split(".")[-1].lower() if "." in file_path else ""
        if ext != "pdf":
            return IngestionService.process_file(file_path)
        
        pages_text = []
        filename = os.path.basename(file_path)
        try:
            with fitz.open(file_path) as doc:
                if doc.is_encrypted:
                    try:
                        doc.authenticate("")
                    except Exception:
                        pass
                for page_num, page in enumerate(doc, start=1):
                    p_text = page.get_text().strip()
                    if p_text:
                        pages_text.append(f"--- PAGE {page_num} ---\n{p_text}")
        except Exception as e:
            print(f"extract_text_with_pages notice for {filename}: {e}")
            return IngestionService.process_file(file_path)

        if not pages_text:
            return IngestionService.process_file(file_path)

        return "\n\n".join(pages_text)

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])

    @staticmethod
    def extract_text_from_pptx(file_path: str) -> str:
        prs = Presentation(file_path)
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
        return "\n".join(text)

    @staticmethod
    def extract_text_from_csv(file_path: str) -> str:
        df = pd.read_csv(file_path)
        return df.to_string()

    @staticmethod
    def extract_text_from_image(file_path: str) -> str:
        try:
            from PIL import Image
            import pytesseract
            img = Image.open(file_path)
            return pytesseract.image_to_string(img)
        except ImportError:
            return "OCR error: PIL or pytesseract not installed."
        except Exception as e:
            return f"OCR error: {str(e)}. Please ensure Tesseract OCR is installed on the system."

    @classmethod
    def process_file(cls, file_path: str) -> str:
        filename = os.path.basename(file_path)
        ext = filename.split(".")[-1].lower() if "." in filename else ""
        try:
            if ext == "pdf":
                return cls.extract_text_from_pdf(file_path)
            elif ext == "docx":
                return cls.extract_text_from_docx(file_path)
            elif ext == "pptx":
                return cls.extract_text_from_pptx(file_path)
            elif ext == "csv":
                return cls.extract_text_from_csv(file_path)
            elif ext in ["xlsx", "xls"]:
                try:
                    df = pd.read_excel(file_path)
                    return df.to_string()
                except Exception:
                    return f"[Spreadsheet Data File: {filename}]"
            elif ext in ["png", "jpg", "jpeg", "webp"]:
                ocr_text = cls.extract_text_from_image(file_path)
                if ocr_text and "OCR error" not in ocr_text:
                    return ocr_text
                return f"[Image File: {filename} ({ext.upper()})]"
            elif ext in ["txt", "md", "json", "py", "html", "css", "js"]:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
            elif ext in ["mp4", "mov", "webm", "avi", "mkv"]:
                return f"[Media Video File Asset: {filename} ({ext.upper()})]"
            elif ext in ["mp3", "wav", "m4a", "ogg", "flac"]:
                return f"[Media Audio File Asset: {filename} ({ext.upper()})]"
            elif ext in ["zip", "rar", "7z", "tar", "gz"]:
                try:
                    import zipfile
                    if ext == "zip" and zipfile.is_zipfile(file_path):
                        with zipfile.ZipFile(file_path, 'r') as z:
                            file_list = z.namelist()
                            return f"[Zip Archive File: {filename}]\nContained Files:\n" + "\n".join(file_list[:30])
                except Exception:
                    pass
                return f"[Compressed Archive File: {filename} ({ext.upper()})]"
            else:
                return f"[Document Asset File: {filename} ({ext.upper()})]"
        except Exception as e:
            return f"[Uploaded Asset File: {filename} - Extraction note: {str(e)}]"

ingestion_service = IngestionService()
